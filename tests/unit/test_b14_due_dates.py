"""B-14: bid due dates from AwardedBidsSearchList.xlsx (fix/b14-bid-date, pending Port sign-off)."""

from __future__ import annotations

from datetime import date

import openpyxl

from bidtab_pipeline.config_loader.silver_config import load_due_dates

HEADER = ["Posted", "Invitation #", "Project Title", "Due Date", "Bidder", "Award Amount"]


def write_export(path, rows, header=HEADER):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, v in enumerate(header, start=1):
        ws.cell(row=5, column=c, value=v)  # PlanetBids export: header on row 5 (read_excel header=4)
    for r, row in enumerate(rows, start=6):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)
    return str(path)


def test_one_date_per_solicitation_latest_wins_blanks_dropped(tmp_path):
    path = write_export(
        tmp_path / "a.xlsx",
        [
            ["06/11/2024", "2024-10553", "P", "08/09/2024 3:00:00 PM", "A", None],
            ["06/11/2024", "2024-10553", "P", "08/09/2024 3:00:00 PM", "B", 100.0],
            ["01/01/2017", " 2017-7224 ", "Q", "03/03/2017 3:00:00 PM", "C", None],
            ["01/01/2017", "2017-7224", "Q", "04/25/2017 3:00:00 PM", "D", 5.0],
            ["01/01/2025", "2024-10691", "R", None, "E", 7.0],
            ["01/01/2025", "2024-9999", "S", "not a date", "F", 7.0],
        ],
    )
    got = dict(load_due_dates(path).itertuples(index=False, name=None))
    assert got == {"2024-10553": date(2024, 8, 9), "2017-7224": date(2017, 4, 25)}


def test_export_without_due_date_column_yields_no_rows(tmp_path):
    path = write_export(tmp_path / "b.xlsx", [["2024-1", "B", 1.0]], header=["Invitation #", "Bidder", "Award Amount"])
    assert load_due_dates(path).empty
