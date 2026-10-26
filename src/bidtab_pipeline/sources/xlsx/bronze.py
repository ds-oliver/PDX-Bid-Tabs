"""Bronze build for the workbook source: port of reference notebook 01's run cells.

Lists every .xlsx at the top level of the source_files volume, parses the
canonical bid tab sheets, and overwrites bronze_raw_rows. Printed output
matches the reference line for line; it is the job log staff read.
"""

from __future__ import annotations

from typing import Any

import openpyxl
from pyspark.sql import SparkSession
from pyspark.sql.functions import current_timestamp
from pyspark.sql.types import (
    BooleanType,
    DateType,
    DoubleType,
    IntegerType,
    StringType,
    StructField,
    StructType,
)

from bidtab_pipeline.io.volumes import list_entry_names
from bidtab_pipeline.settings import Settings
from bidtab_pipeline.sources.xlsx.parser import parse_mapped_sheet_rows

TARGET_TABLE_NAME = "bronze_raw_rows"
SOURCE_VOLUME = "source_files"

# Hard-coded in the baseline; kept as code, not config (docs/backlog.md B-04).
SKIP_FILES = [
    "PDX_2017D013.xlsx",
]

# Canonical sheet names the parser will look for in each workbook
CANONICAL_SHEETS = [
    "Base Bid Tab",
    "Alternate Bid Tab",
    "Alternate Bid Tab 1",
    "Alternate Bid Tab 2",
    "Alternate Bid Tab 3",
    "Alternate Bid Tab 4",
]

BRONZE_SCHEMA = StructType(
    [
        StructField("source_file", StringType(), True),
        StructField("source_sheet", StringType(), True),
        StructField("bid_type", StringType(), True),
        StructField("ean", StringType(), True),
        StructField("project_name", StringType(), True),
        StructField("airport_name", StringType(), True),
        StructField("solicitation_no", StringType(), True),
        StructField("bid_date", DateType(), True),
        StructField("item_no", StringType(), True),
        StructField("item_description_raw", StringType(), True),
        StructField("estimated_quantity", DoubleType(), True),
        StructField("unit", StringType(), True),
        StructField("contractor_name", StringType(), True),
        StructField("contractor_position", IntegerType(), True),
        StructField("is_engineers_estimate", BooleanType(), True),
        StructField("unit_price", DoubleType(), True),
        StructField("total_price", DoubleType(), True),
        StructField("is_section_header", BooleanType(), True),
    ]
)


def list_top_level_xlsx(volume_path: str, dbutils: Any = None) -> list[str]:
    files = []
    for name in list_entry_names(volume_path, dbutils):
        if name.endswith("/"):
            continue
        if name.lower().endswith(".xlsx") and not name.startswith("~$"):
            files.append(name)
    return sorted(files)


def parse_workbook_auto(volume_path: str, source_file: str) -> tuple[list[dict[str, Any]], int, list[str]]:
    """Parse a workbook by auto-detecting canonical sheet names."""
    workbook_path = f"{volume_path}/{source_file}"
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    rows_out: list[dict[str, Any]] = []
    parsed_sheet_count = 0
    warnings: list[str] = []

    # Build case-insensitive index of actual sheet names
    sheet_name_index = {name.lower().strip(): name for name in wb.sheetnames}

    for canonical_name in CANONICAL_SHEETS:
        key = canonical_name.lower().strip()
        if key not in sheet_name_index:
            continue

        workbook_sheet_name = sheet_name_index[key]

        ws = wb[workbook_sheet_name]
        sheet_rows = parse_mapped_sheet_rows(
            ws=ws,
            source_file=source_file,
            source_sheet=workbook_sheet_name,
            bid_type=canonical_name,
        )
        if sheet_rows:
            parsed_sheet_count += 1
            rows_out.extend(sheet_rows)

    if parsed_sheet_count == 0:
        warnings.append(f"no_canonical_sheets_found:available={wb.sheetnames}")

    return rows_out, parsed_sheet_count, warnings


def build_bronze(spark: SparkSession, settings: Settings, dbutils: Any = None) -> dict[str, Any]:
    """Bulk ingest: parse every top-level workbook and overwrite bronze_raw_rows.

    Returns the run summary (files seen, skipped, errored, rows written).
    """
    settings.require_table_write()
    volume_path = settings.volume(SOURCE_VOLUME)
    target_table = settings.table(TARGET_TABLE_NAME)

    files_seen = list_top_level_xlsx(volume_path, dbutils)
    skipped_files: list[str] = []
    errored_files: list[tuple[str, str]] = []
    file_summaries: list[dict[str, Any]] = []
    bulk_rows: list[dict[str, Any]] = []

    for filename in files_seen:
        if filename in SKIP_FILES:
            skipped_files.append(filename)
            continue
        try:
            rows_out, parsed_sheets, warnings = parse_workbook_auto(volume_path, filename)
            bulk_rows.extend(rows_out)
            file_summaries.append(
                {
                    "filename": filename,
                    "sheets_parsed": parsed_sheets,
                    "rows_emitted": len(rows_out),
                    "warnings": "; ".join(warnings),
                }
            )
        except Exception as exc:
            errored_files.append((filename, str(exc)))
            file_summaries.append(
                {
                    "filename": filename,
                    "sheets_parsed": 0,
                    "rows_emitted": 0,
                    "warnings": f"error:{exc}",
                }
            )

    rows_written = 0
    if bulk_rows:
        df_spark = spark.createDataFrame(bulk_rows, schema=BRONZE_SCHEMA).withColumn(
            "bronze_loaded_at", current_timestamp()
        )
        (df_spark.write.format("delta").mode("overwrite").option("mergeSchema", "true").saveAsTable(target_table))
        rows_written = len(bulk_rows)

    print("filename | sheets_parsed | rows_emitted | warnings")
    for s in file_summaries:
        print(f"{s['filename']} | {s['sheets_parsed']} | {s['rows_emitted']} | {s['warnings']}")

    processed_count = len([f for f in files_seen if f not in SKIP_FILES])
    print("\nFINAL SUMMARY")
    print(f"total_files_seen={len(files_seen)}")
    print(f"total_files_processed={processed_count}")
    print(f"total_files_skipped={len(skipped_files)}")
    print(f"total_files_errored={len(errored_files)}")
    print(f"total_rows_written={rows_written}")
    if skipped_files:
        print(f"skipped_files={skipped_files}")
    if errored_files:
        print("errored_files:")
        for filename, err in errored_files:
            print(f"  - {filename}: {err}")

    return {
        "files_seen": files_seen,
        "skipped_files": skipped_files,
        "errored_files": errored_files,
        "file_summaries": file_summaries,
        "rows_written": rows_written,
    }
