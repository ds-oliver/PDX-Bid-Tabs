# Databricks notebook source
from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from typing import Any

import openpyxl
import pandas as pd
from pyspark.sql import SparkSession
from pyspark.sql.functions import current_timestamp
from pyspark.sql.types import (
    BooleanType,
    DateType,
    DoubleType,
    IntegerType,
    StringType,
    StructField,
    StructType,
)


# COMMAND ----------

# Config
VOLUME_PATH = "/Volumes/pop-poc/bidtab/source_files"
SKIP_FILES = [
    "PDX_2017D013.xlsx",
]
TARGET_TABLE = "`pop-poc`.`bidtab`.`bronze_raw_rows`"

# Canonical sheet names the parser will look for in each workbook
CANONICAL_SHEETS = [
    "Base Bid Tab",
    "Alternate Bid Tab",
    "Alternate Bid Tab 1",
    "Alternate Bid Tab 2",
    "Alternate Bid Tab 3",
    "Alternate Bid Tab 4",
]

# COMMAND ----------

# Standalone parser helpers (mirrors current local pipeline behavior)
WS_RE = re.compile(r"\s+")
TOTALS_PATTERN = re.compile(r"(total amount|basis of award)", re.IGNORECASE)
EAN_RE = re.compile(r"EAN\s+([A-Za-z0-9\-]+)", re.IGNORECASE)
SOL_RE = re.compile(r"Solicitation\s+No\.\s*[:\-]?\s*([\d]+-[\d]+)", re.IGNORECASE)
DATE_RE = re.compile(
    r"(\d{1,2}[\-/]\d{1,2}[\-/]\d{2,4}|[A-Za-z]{3,9}\s+\d{1,2},\s*\d{4}|\d{4}[\-/]\d{1,2}[\-/]\d{1,2})"
)
CANONICAL_BID_TYPE_RE = re.compile(r"^(Base Bid Tab|Alternate Bid Tab(?: \d+)?)$", re.IGNORECASE)


@dataclass
class BidderBlock:
    bidder_name_raw: str
    unit_price_col: int
    total_price_col: int


@dataclass
class TableParseResult:
    table_header_row: int
    col_map: dict[str, int]
    bidder_blocks: list[BidderBlock]
    warnings: list[str]


def clean_whitespace(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ").strip()
    return WS_RE.sub(" ", text)


def normalize_text(value: Any) -> str:
    return clean_whitespace(value)


def parse_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = clean_whitespace(value)
    if not text or text == "#REF!":
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.replace("$", "").replace(",", "").replace("(", "").replace(")", "")
    if text in {"", "-", "--"}:
        return None
    try:
        dec = Decimal(text)
    except InvalidOperation:
        return None
    return -dec if negative else dec


def to_float(value: Any) -> float | None:
    dec = parse_decimal(value)
    return float(dec) if dec is not None else None


def classify_bidder(name: str) -> tuple[str, bool]:
    lower = clean_whitespace(name).lower()
    is_est = "engineer" in lower and "estimate" in lower
    bidder_type = "ENGINEERS_ESTIMATE" if is_est else "CONTRACTOR"
    return bidder_type, is_est


def parse_bid_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_whitespace(value)
    if not text:
        return None
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def _parse_date_iso(value: str) -> str:
    if not value:
        return ""
    for fmt in ["%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%b %d, %Y", "%B %d, %Y"]:
        try:
            return datetime.strptime(value.strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _norm(value: Any) -> str:
    return clean_whitespace(value).lower()


def detect_table_header_row(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[int | None, str]:
    candidates: list[int] = []
    for r in range(1, ws.max_row + 1):
        row_vals = [_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        if any(v in {"item no.", "item no"} for v in row_vals):
            if row_vals.count("unit price") >= 2 and row_vals.count("total price") >= 2:
                candidates.append(r)
    if not candidates:
        return None, "missing_item_no"

    scored: list[tuple[int, int]] = []
    required = {"item no.", "item description", "estimated quantity", "units"}
    for r in candidates:
        row_vals = [_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        score = sum(1 for v in required if v in row_vals)
        scored.append((score, r))
    scored.sort(reverse=True)
    best_score, best_row = scored[0]
    tied = [row for score, row in scored if score == best_score]
    if len(tied) > 1:
        return None, "ambiguous_header_rows"
    return best_row, "ok"


def map_base_columns(ws: openpyxl.worksheet.worksheet.Worksheet, table_header_row: int) -> dict[str, int]:
    col_map: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = _norm(ws.cell(row=table_header_row, column=c).value)
        if v in {"item no.", "item no"}:
            col_map["line_no_col"] = c
        elif v == "item description":
            col_map["desc_col"] = c
        elif v == "estimated quantity":
            col_map["qty_col"] = c
        elif v == "units":
            col_map["unit_col"] = c
    return col_map


def detect_bidder_blocks(ws: openpyxl.worksheet.worksheet.Worksheet, table_header_row: int) -> tuple[list[BidderBlock], list[str]]:
    unit_cols: list[int] = []
    total_cols: list[int] = []
    for c in range(1, ws.max_column + 1):
        v = _norm(ws.cell(row=table_header_row, column=c).value)
        if v == "unit price":
            unit_cols.append(c)
        elif v == "total price":
            total_cols.append(c)
    warnings: list[str] = []
    if len(unit_cols) != len(total_cols):
        warnings.append("bidder_pair_count_mismatch")
    blocks: list[BidderBlock] = []
    for i in range(min(len(unit_cols), len(total_cols))):
        u = unit_cols[i]
        t = total_cols[i]
        bidder_name = "UNKNOWN_BIDDER"
        for r in range(table_header_row - 1, max(0, table_header_row - 10) - 1, -1):
            v_unit = clean_whitespace(ws.cell(row=r, column=u).value)
            v_total = clean_whitespace(ws.cell(row=r, column=t).value)
            if v_unit and v_total and v_unit == v_total:
                bidder_name = v_unit
                break
            if v_unit and not v_total:
                bidder_name = v_unit
                break
            if v_total and not v_unit:
                bidder_name = v_total
                break
        blocks.append(BidderBlock(bidder_name_raw=bidder_name, unit_price_col=u, total_price_col=t))
    return blocks, warnings


def parse_table_structure(ws: openpyxl.worksheet.worksheet.Worksheet) -> TableParseResult:
    row, status = detect_table_header_row(ws)
    if row is None:
        return TableParseResult(table_header_row=-1, col_map={}, bidder_blocks=[], warnings=[status])
    col_map = map_base_columns(ws, row)
    warnings: list[str] = []
    missing = [k for k in ["line_no_col", "desc_col", "qty_col", "unit_col"] if k not in col_map]
    if missing:
        warnings.append("missing_base_columns:" + ",".join(missing))
    bidder_blocks, bidder_warnings = detect_bidder_blocks(ws, row)
    warnings.extend(bidder_warnings)
    if not bidder_blocks:
        warnings.append("no_bidder_blocks")
    return TableParseResult(table_header_row=row, col_map=col_map, bidder_blocks=bidder_blocks, warnings=warnings)


def parse_project_header(ws: openpyxl.worksheet.worksheet.Worksheet, table_header_row: int) -> dict[str, str]:
    header_lines: list[str] = []
    max_scan_col = min(ws.max_column, 20)
    for r in range(1, table_header_row):
        vals = [clean_whitespace(ws.cell(row=r, column=c).value) for c in range(1, max_scan_col + 1)]
        vals = [v for v in vals if v]
        if vals:
            header_lines.append(" | ".join(vals))
    joined = "\n".join(header_lines)
    ean = clean_whitespace(EAN_RE.search(joined).group(1)) if EAN_RE.search(joined) else ""
    sol = clean_whitespace(SOL_RE.search(joined).group(1)) if SOL_RE.search(joined) else ""
    date_match = DATE_RE.search(joined)
    letting_raw = clean_whitespace(date_match.group(1)) if date_match else ""
    return {
        "project_ean": ean,
        "solicitation_no": sol,
        "letting_date_raw": letting_raw,
        "letting_date": _parse_date_iso(letting_raw),
        "location_name_raw": clean_whitespace(ws.cell(row=1, column=1).value),
        "project_name_raw": clean_whitespace(ws.cell(row=2, column=1).value),
    }


def _extract_line_no(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int) -> str:
    cell = ws.cell(row=row, column=col)
    raw = cell.value
    if raw is None:
        return ""
    if isinstance(raw, str):
        return clean_whitespace(raw)
    if isinstance(raw, (int, float)):
        num_format = clean_whitespace(cell.number_format)
        if num_format and set(num_format) <= {"0"}:
            try:
                return str(int(raw)).zfill(len(num_format))
            except Exception:
                pass
        if int(raw) == raw:
            return str(int(raw))
    return clean_whitespace(raw)


def iter_item_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    table_header_row: int,
    col_map: dict[str, int],
    bidder_blocks: list[BidderBlock],
    termination_blank_streak: int = 10,
):
    blank_streak = 0
    row = table_header_row + 1
    while row <= ws.max_row:
        line_no = _extract_line_no(ws, row, col_map["line_no_col"])
        desc = clean_whitespace(ws.cell(row=row, column=col_map["desc_col"]).value)
        qty = parse_decimal(ws.cell(row=row, column=col_map["qty_col"]).value)
        unit_raw = clean_whitespace(ws.cell(row=row, column=col_map["unit_col"]).value)

        if not line_no and not desc:
            blank_streak += 1
            if blank_streak >= termination_blank_streak:
                break
            row += 1
            continue
        blank_streak = 0
        is_totals = bool(TOTALS_PATTERN.search(desc))
        if line_no or is_totals:
            for block in bidder_blocks:
                unit_price = parse_decimal(ws.cell(row=row, column=block.unit_price_col).value)
                total_price = parse_decimal(ws.cell(row=row, column=block.total_price_col).value)
                yield {
                    "line_no": "" if is_totals else line_no,
                    "item_description_raw": desc,
                    "quantity": None if is_totals else (float(qty) if qty is not None else None),
                    "unit_code_raw": "" if is_totals else unit_raw,
                    "unit_price": None if is_totals else (float(unit_price) if unit_price is not None else None),
                    "total_price": None if is_totals else (float(total_price) if total_price is not None else None),
                    "is_totals_row": is_totals,
                    "bidder_name_raw": block.bidder_name_raw,
                }
        row += 1


def extract_sheet_metadata(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    structure = parse_table_structure(ws)
    if structure.table_header_row < 0:
        return {
            "airport_name": normalize_text(ws.cell(row=1, column=1).value),
            "project_name": normalize_text(ws.cell(row=2, column=1).value),
            "ean": "",
            "solicitation_no": "",
            "bid_date": parse_bid_date(ws.cell(row=6, column=1).value),
            "table_header_row": -1,
            "structure": structure,
        }

    header = parse_project_header(ws, structure.table_header_row)
    bid_date = parse_bid_date(header.get("letting_date", "")) or parse_bid_date(ws.cell(row=6, column=1).value)

    return {
        "airport_name": header.get("location_name_raw", "") or normalize_text(ws.cell(row=1, column=1).value),
        "project_name": header.get("project_name_raw", "") or normalize_text(ws.cell(row=2, column=1).value),
        "ean": header.get("project_ean", ""),
        "solicitation_no": header.get("solicitation_no", ""),
        "bid_date": bid_date,
        "table_header_row": structure.table_header_row,
        "structure": structure,
    }


def parse_mapped_sheet_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    source_file: str,
    source_sheet: str,
    bid_type: str,
) -> list[dict[str, Any]]:
    meta = extract_sheet_metadata(ws)
    structure = meta["structure"]
    if structure.table_header_row < 0 or not structure.bidder_blocks:
        print(f"WARNING: unable to parse bidder matrix for mapped sheet '{source_sheet}', skipping.")
        return []
    if structure.warnings:
        print(f"WARNING: parse warnings for sheet '{source_sheet}': {';'.join(structure.warnings)}")

    bidder_position_by_name: dict[str, int] = {}
    for idx, block in enumerate(structure.bidder_blocks, start=1):
        bidder_name = normalize_text(block.bidder_name_raw)
        if bidder_name and bidder_name not in bidder_position_by_name:
            bidder_position_by_name[bidder_name] = idx

    rows: list[dict[str, Any]] = []
    for item in iter_item_rows(
        ws,
        structure.table_header_row,
        structure.col_map,
        structure.bidder_blocks,
    ):
        if item.get("is_totals_row"):
            continue
        bidder_raw = normalize_text(item.get("bidder_name_raw", ""))
        _, is_engineers_estimate = classify_bidder(bidder_raw)
        # EE rows are retained (not filtered) so downstream consumers can compute
        # variance/baseline analyses. The is_engineers_estimate flag distinguishes
        # them from contractor bids.

        item_no = normalize_text(item.get("line_no", ""))
        desc = normalize_text(item.get("item_description_raw", ""))
        contractor_name = bidder_raw
        if not contractor_name:
            fallback_pos = len(bidder_position_by_name) + 1
            contractor_name = f"Name of Contractor {fallback_pos}"

        contractor_position = bidder_position_by_name.get(bidder_raw)
        if contractor_position is None:
            contractor_position = len(bidder_position_by_name) + 1

        rows.append(
            {
                "source_file": source_file,
                "source_sheet": source_sheet,
                "bid_type": bid_type,
                "ean": meta["ean"],
                "project_name": meta["project_name"],
                "airport_name": meta["airport_name"],
                "solicitation_no": meta["solicitation_no"],
                "bid_date": meta["bid_date"],
                "item_no": item_no if item_no else None,
                "item_description_raw": desc,
                "estimated_quantity": to_float(item.get("quantity")),
                "unit": normalize_text(item.get("unit_code_raw", "")) or None,
                "contractor_name": contractor_name,
                "contractor_position": int(contractor_position),
                "is_engineers_estimate": bool(is_engineers_estimate),
                "unit_price": to_float(item.get("unit_price")),
                "total_price": to_float(item.get("total_price")),
                "is_section_header": bool((not item_no) and bool(desc)),
            }
        )
    return rows


# COMMAND ----------

# Common execution helpers
def list_top_level_xlsx(volume_path: str) -> list[str]:
    files = []
    for info in dbutils.fs.ls(volume_path):
        name = str(info.name)
        if name.endswith("/"):
            continue
        if name.lower().endswith(".xlsx") and not name.startswith("~$"):
            files.append(name)
    return sorted(files)


def parse_workbook_auto(source_file: str) -> tuple[list[dict[str, Any]], int, list[str]]:
    """Parse a workbook by auto-detecting canonical sheet names."""
    workbook_path = f"{VOLUME_PATH}/{source_file}"
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    rows_out: list[dict[str, Any]] = []
    parsed_sheet_count = 0
    warnings: list[str] = []

    # Build case-insensitive index of actual sheet names
    sheet_name_index = {name.lower().strip(): name for name in wb.sheetnames}

    for canonical_name in CANONICAL_SHEETS:
        key = canonical_name.lower().strip()
        if key not in sheet_name_index:
            continue

        workbook_sheet_name = sheet_name_index[key]

        ws = wb[workbook_sheet_name]
        sheet_rows = parse_mapped_sheet_rows(
            ws=ws,
            source_file=source_file,
            source_sheet=workbook_sheet_name,
            bid_type=canonical_name,
        )
        if sheet_rows:
            parsed_sheet_count += 1
            rows_out.extend(sheet_rows)

    if parsed_sheet_count == 0:
        warnings.append(f"no_canonical_sheets_found:available={wb.sheetnames}")

    return rows_out, parsed_sheet_count, warnings


spark = SparkSession.builder.getOrCreate()
schema = StructType(
    [
        StructField("source_file", StringType(), True),
        StructField("source_sheet", StringType(), True),
        StructField("bid_type", StringType(), True),
        StructField("ean", StringType(), True),
        StructField("project_name", StringType(), True),
        StructField("airport_name", StringType(), True),
        StructField("solicitation_no", StringType(), True),
        StructField("bid_date", DateType(), True),
        StructField("item_no", StringType(), True),
        StructField("item_description_raw", StringType(), True),
        StructField("estimated_quantity", DoubleType(), True),
        StructField("unit", StringType(), True),
        StructField("contractor_name", StringType(), True),
        StructField("contractor_position", IntegerType(), True),
        StructField("is_engineers_estimate", BooleanType(), True),
        StructField("unit_price", DoubleType(), True),
        StructField("total_price", DoubleType(), True),
        StructField("is_section_header", BooleanType(), True),
    ]
)

# Bulk ingest — all files
files_seen = list_top_level_xlsx(VOLUME_PATH)
skipped_files: list[str] = []
errored_files: list[tuple[str, str]] = []
file_summaries: list[dict[str, Any]] = []
bulk_rows: list[dict[str, Any]] = []

for filename in files_seen:
    if filename in SKIP_FILES:
        skipped_files.append(filename)
        continue
    try:
        rows_out, parsed_sheets, warnings = parse_workbook_auto(filename)
        bulk_rows.extend(rows_out)
        file_summaries.append(
            {
                "filename": filename,
                "sheets_parsed": parsed_sheets,
                "rows_emitted": len(rows_out),
                "warnings": "; ".join(warnings),
            }
        )
    except Exception as exc:
        errored_files.append((filename, str(exc)))
        file_summaries.append(
            {
                "filename": filename,
                "sheets_parsed": 0,
                "rows_emitted": 0,
                "warnings": f"error:{exc}",
            }
        )

rows_written = 0
if bulk_rows:
    df_spark = spark.createDataFrame(bulk_rows, schema=schema).withColumn("bronze_loaded_at", current_timestamp())
    (
        df_spark.write.format("delta")
        .mode("overwrite")
        .option("mergeSchema", "true")
        .saveAsTable(TARGET_TABLE)
    )
    rows_written = len(bulk_rows)

print("filename | sheets_parsed | rows_emitted | warnings")
for s in file_summaries:
    print(f"{s['filename']} | {s['sheets_parsed']} | {s['rows_emitted']} | {s['warnings']}")

processed_count = len([f for f in files_seen if f not in SKIP_FILES])
print("\nFINAL SUMMARY")
print(f"total_files_seen={len(files_seen)}")
print(f"total_files_processed={processed_count}")
print(f"total_files_skipped={len(skipped_files)}")
print(f"total_files_errored={len(errored_files)}")
print(f"total_rows_written={rows_written}")
if skipped_files:
    print(f"skipped_files={skipped_files}")
if errored_files:
    print("errored_files:")
    for filename, err in errored_files:
        print(f"  - {filename}: {err}")